import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "cf-kit"))

try:
    import openpyxl
    import build_cfkit
except ImportError:  # openpyxl or litsa_design not installed
    openpyxl = None


@unittest.skipIf(openpyxl is None, "pip install -r requirements.txt to run the cf-kit tests")
class CfKit(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.path = build_cfkit.build(Path(tempfile.mkdtemp()) / "cf-kit.xlsx")
        cls.wb = openpyxl.load_workbook(cls.path)

    def test_tabs_in_order(self):
        self.assertEqual(self.wb.sheetnames[:2], ["Start Here", "Dashboard"])
        for name in build_cfkit.TABS:
            self.assertIn(name, self.wb.sheetnames)

    def test_gap_formula_compares_with_previous_row(self):
        ws = self.wb["Residences"]
        self.assertIsNone(ws["L2"].value)
        self.assertEqual(ws["L3"].value, '=IF(AND(A3<>"",B2<>""),A3-B2-1,"")')

    def test_dashboard_points_at_real_columns(self):
        ws = self.wb["Dashboard"]
        formulas = " ".join(str(ws.cell(row=r, column=2).value) for r in range(2, 15))
        self.assertIn("Residences!L3", formulas)
        self.assertIn("Employment!M3", formulas)
        self.assertIn("Employment!K2", formulas)
        self.assertEqual(self.wb["Employment"]["K1"].value, "Needs explanation")
        self.assertEqual(self.wb["Employment"]["M1"].value, "Days since previous job")
        self.assertEqual(ws["A13"].value, "Last updated")

    def test_documents_prefilled(self):
        self.assertTrue(self.wb["Documents"]["A2"].value.startswith("Driving record"))

    def test_headers_frozen(self):
        for name in build_cfkit.TABS:
            self.assertEqual(self.wb[name].freeze_panes, "A2")


if __name__ == "__main__":
    unittest.main()
