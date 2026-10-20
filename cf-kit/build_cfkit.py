#!/usr/bin/env python3
"""Generate the character & fitness spreadsheet (cf-kit.xlsx).

    python3 cf-kit/build_cfkit.py            # writes cf-kit/dist/cf-kit.xlsx

The layout is defined in TABS below and described for people in cf-kit/SPEC.md. Colors come from
the LITSA design system (litsa_design). Only features that work in both Excel and Google Sheets
are used: plain formulas, dropdown lists, conditional formatting, frozen header rows.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from litsa_design import tokens

VERSION = "0.1"
ROOT = Path(__file__).resolve().parent
ROWS = 60  # blank, pre-formatted rows per tab
HEX = lambda h: h.lstrip("#").upper()
HEAD_FILL = PatternFill("solid", fgColor=HEX(tokens.BLUE_DEEP))
HEAD_FONT = Font(bold=True, color="FFFFFF")
CALC_FILL = PatternFill("solid", fgColor=HEX(tokens.PAPER))
ALERT_FILL = PatternFill("solid", fgColor="F6D5D0")
WARN_FILL = PatternFill("solid", fgColor="FBEBC5")
DATE = "yyyy-mm-dd"

# Each column: (header, width, kind) where kind is "text", "date", "money", a list of dropdown
# choices, or ("formula", template) with {r} for the row and {p} for the previous row.
TABS = {
    "Residences": [
        ("From", 12, "date"), ("To", 12, "date"), ("Street address", 30, "text"), ("Unit", 8, "text"),
        ("City", 16, "text"), ("State", 8, "text"), ("ZIP", 9, "text"), ("Country", 12, "text"),
        ("Lived with", 18, "text"), ("Landlord or contact", 24, "text"), ("Notes", 30, "text"),
        ("Days since previous address", 14, ("formula", '=IF(AND(A{r}<>"",B{p}<>""),A{r}-B{p}-1,"")')),
    ],
    "Employment": [
        ("From", 12, "date"), ("To", 12, "date"), ("Employer (or \"Unemployed\")", 28, "text"), ("Employer address", 28, "text"),
        ("Supervisor", 20, "text"), ("Supervisor phone", 15, "text"), ("Supervisor email", 24, "text"), ("Title", 20, "text"),
        ("Hours", 14, ["Full-time", "Part-time", "Seasonal", "Unemployed"]),
        ("Reason for leaving", 20, ["Still there", "Graduated or returned to school", "Resigned", "Contract ended", "Laid off", "Asked to resign", "Terminated", "Other"]),
        ("Needs explanation", 12, ("formula", '=IF(OR(J{r}="Terminated",J{r}="Asked to resign"),"Yes","")')),
        ("Explanation or notes", 34, "text"),
        ("Days since previous job", 14, ("formula", '=IF(AND(A{r}<>"",B{p}<>""),A{r}-B{p}-1,"")')),
    ],
    "Education": [
        ("From", 12, "date"), ("To", 12, "date"), ("School", 28, "text"), ("City and state", 18, "text"),
        ("Degree or program", 20, "text"), ("Discipline or probation?", 14, ["No", "Yes"]), ("Explanation", 40, "text"),
    ],
    "Legal Matters": [
        ("Date", 12, "date"),
        ("Type", 22, ["Traffic citation", "Parking (if asked)", "Arrest", "Criminal charge", "Civil case (party)", "Restraining order", "Juvenile matter", "Other"]),
        ("Court or agency", 24, "text"), ("Docket or citation number", 18, "text"), ("What happened", 34, "text"),
        ("Outcome", 22, "text"), ("Sealed, expunged or dismissed?", 14, ["No", "Sealed", "Expunged", "Dismissed", "Continued without a finding"]),
        ("Records obtained?", 12, ["No", "Requested", "Yes"]), ("Notes", 30, "text"),
    ],
    "Financial": [
        ("Item", 24, "text"), ("Type", 18, ["Collections", "Default", "Judgment", "Bankruptcy", "Tax issue", "Student loan", "Other"]),
        ("Amount", 12, "money"), ("Date", 12, "date"), ("Current status", 24, "text"), ("Resolved on", 12, "date"), ("Notes", 30, "text"),
    ],
    "Licenses": [
        ("Type", 22, ["Driver's license", "Professional license", "Other"]), ("State or issuer", 14, "text"),
        ("Number (last 4 only)", 12, "text"), ("Issued", 12, "date"), ("Expires", 12, "date"),
        ("Status", 14, ["Active", "Expired", "Suspended", "Revoked", "Surrendered"]), ("Notes", 30, "text"),
    ],
    "References": [
        ("Name", 22, "text"), ("Relationship", 20, "text"), ("Years known", 10, "text"), ("Phone", 15, "text"),
        ("Email", 26, "text"), ("Asked on", 12, "date"), ("Agreed?", 10, ["Yes", "No", "Waiting"]),
    ],
    "Law School App Check": [
        ("Question on your law school application", 40, "text"), ("What you answered then", 28, "text"),
        ("What your records show now", 28, "text"), ("Match or amend?", 14, ["Match", "Amend"]),
        ("Amendment filed on", 14, "date"), ("Notes", 30, "text"),
    ],
    "Documents": [
        ("Document", 34, "text"), ("Where to get it", 36, "text"), ("Requested on", 13, "date"), ("Received on", 13, "date"),
        ("Status", 13, ("formula", '=IF(D{r}<>"","Received",IF(AND(C{r}<>"",TODAY()-C{r}>30),"Overdue",IF(C{r}<>"","Requested","")))')),
        ("Notes", 30, "text"),
    ],
    "Change Log": [("Date", 12, "date"), ("What changed", 60, "text")],
}

DOCUMENTS = [
    ("Driving record, for every state where you've held a license", "Each state's motor vehicle agency (in Massachusetts, the RMV on mass.gov)"),
    ("Credit report", "annualcreditreport.com, the free official site for all three bureaus"),
    ("Court records for each matter on the Legal Matters tab", "The clerk's office of the court that handled it"),
    ("A copy of your law school application", "Suffolk Law Admissions or the Dean of Students office"),
    ("Transcripts and any discipline records from each school", "Each school's registrar or dean of students"),
]

START = [
    ("Character & fitness kit", "title"),
    (f"Version {VERSION}, from LITSA at Suffolk Law. github.com/suffolklitsa/suffolk-guides", "small"),
    ("", None),
    ("What this is", "head"),
    ("Bar applications ask for every address, job, school, legal matter and financial problem going back years. Filling this in as you go means you won't be rebuilding a decade from memory in 3L year.", None),
    ("", None),
    ("Your information stays with you", "head"),
    ("This file lives wherever you save it. LITSA never sees it. Don't email it around or upload it to shared folders; keep it somewhere private and backed up.", None),
    ("", None),
    ("How to use it", "head"),
    ("1. Fill in each tab, oldest entries first. Gray columns calculate themselves; don't type in them.", None),
    ("2. Check the Dashboard. Red numbers are gaps or items that need an explanation.", None),
    ("3. On Law School App Check, compare your law school application answers with your records. Anything marked Amend goes to the Dean of Students before you file for the bar.", None),
    ("4. Request documents early on the Documents tab; requests over 30 days old turn amber.", None),
    ("5. Add a line to the Change Log each time you update. Review every semester.", None),
    ("", None),
    ("This is not legal advice", "head"),
    ("Each jurisdiction asks different questions. Always answer the questions on the actual application, completely and truthfully, and ask the Dean of Students or the bar examiners if you're unsure whether something must be disclosed.", None),
]


def style_header(ws, columns):
    for c, (name, width, _) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"


def build_tab(wb, name, columns):
    ws = wb.create_sheet(name)
    style_header(ws, columns)
    last = ROWS + 1
    for c, (_, _, kind) in enumerate(columns, 1):
        col = get_column_letter(c)
        rng = f"{col}2:{col}{last}"
        if kind == "date":
            for r in range(2, last + 1):
                ws.cell(row=r, column=c).number_format = DATE
            dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(1950,1,1)", allow_blank=True,
                                error="Enter a date, for example 2024-09-01.", errorTitle="Date")
            ws.add_data_validation(dv)
            dv.add(rng)
        elif kind == "money":
            for r in range(2, last + 1):
                ws.cell(row=r, column=c).number_format = '"$"#,##0.00'
        elif isinstance(kind, list):
            dv = DataValidation(type="list", formula1='"' + ",".join(kind) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(rng)
        elif isinstance(kind, tuple) and kind[0] == "formula":
            for r in range(2, last + 1):
                if "{p}" in kind[1] and r == 2:
                    continue
                cell = ws.cell(row=r, column=c, value=kind[1].format(r=r, p=r - 1))
                cell.fill = CALC_FILL
            if name in ("Residences", "Employment") and "Days since" in columns[c - 1][0]:
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["30"], fill=ALERT_FILL))
            if columns[c - 1][0] == "Needs explanation":
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Yes"'], fill=ALERT_FILL))
            if columns[c - 1][0] == "Status" and name == "Documents":
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Overdue"'], fill=WARN_FILL))
        if columns[c - 1][1] >= 28:
            for r in range(2, last + 1):
                ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    if name == "Law School App Check":
        ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="equal", formula=['"Amend"'], fill=ALERT_FILL))
    return ws


def dashboard(wb):
    ws = wb.create_sheet("Dashboard", 1)
    n = ROWS + 1
    rows = [
        ("Addresses entered", f"=COUNTA(Residences!A2:A{n})", None),
        ("Address gaps over 30 days", f"=COUNTIF(Residences!L3:L{n},\">30\")", "alert"),
        ("Jobs and gaps entered", f"=COUNTA(Employment!A2:A{n})", None),
        ("Employment gaps over 30 days", f"=COUNTIF(Employment!M3:M{n},\">30\")", "alert"),
        ("Jobs that need an explanation", f"=COUNTIF(Employment!K2:K{n},\"Yes\")", "alert"),
        ("Schools entered", f"=COUNTA(Education!A2:A{n})", None),
        ("Legal matters entered", f"=COUNTA('Legal Matters'!A2:A{n})", None),
        ("Financial items not yet resolved", f"=COUNTIFS(Financial!A2:A{n},\"<>\",Financial!F2:F{n},\"\")", "alert"),
        ("References who agreed", f"=COUNTIF(References!G2:G{n},\"Yes\")", None),
        ("Law school answers to amend", f"=COUNTIF('Law School App Check'!D2:D{n},\"Amend\")", "alert"),
        ("Documents still outstanding", f"=COUNTIF(Documents!E2:E{n},\"Requested\")+COUNTIF(Documents!E2:E{n},\"Overdue\")", "warn"),
        ("Last updated", f"=IF(COUNT('Change Log'!A2:A{n})=0,\"\",MAX('Change Log'!A2:A{n}))", "date"),
        ("Next review due", "=IF(B13=\"\",\"\",B13+120)", "date"),
    ]
    ws["A1"], ws["B1"] = "Where things stand", ""
    ws["A1"].font = Font(bold=True, size=14, color=HEX(tokens.INK))
    for i, (label, formula, kind) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=label)
        cell = ws.cell(row=i, column=2, value=formula)
        cell.fill = CALC_FILL
        if kind == "date":
            cell.number_format = DATE
        elif kind in ("alert", "warn"):
            ws.conditional_formatting.add(f"B{i}", CellIsRule(operator="greaterThan", formula=["0"], fill=ALERT_FILL if kind == "alert" else WARN_FILL))
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 14


def start_here(wb):
    ws = wb.active
    ws.title = "Start Here"
    ws.column_dimensions["A"].width = 100
    for r, (text, kind) in enumerate(START, 1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            cell.font = Font(bold=True, size=18, color=HEX(tokens.INK))
        elif kind == "head":
            cell.font = Font(bold=True, size=12, color=HEX(tokens.BLUE_DEEP))
        elif kind == "small":
            cell.font = Font(size=9, color=HEX(tokens.INK_3))


def build(path: Path) -> Path:
    wb = Workbook()
    start_here(wb)
    for name, columns in TABS.items():
        ws = build_tab(wb, name, columns)
        if name == "Documents":
            for r, (doc, where) in enumerate(DOCUMENTS, 2):
                ws.cell(row=r, column=1, value=doc)
                ws.cell(row=r, column=2, value=where)
    dashboard(wb)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = HEX(tokens.GOLD if ws.title in ("Start Here", "Dashboard") else tokens.BLUE)
    wb.properties.title = "Character & fitness kit"
    wb.properties.creator = "LITSA, Suffolk Law"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    print(f"wrote {build(ROOT / 'dist' / 'cf-kit.xlsx')}")
